# /// script
# requires-python = ">=3.10"
# dependencies = ["python-pptx>=1.0", "pillow>=10"]
# ///
"""Build the mid-semester deck from the college template.

Run from the repo root:  uv run presentation/build_deck.py
Real plots are taken from figures/<name>.png (made by analysis/make_figures.py). CST screenshots go in
figures/cst_*.png. Any missing file becomes a marked "Result pending" box, so re-run after each CST step.
Fill NUMBERS from figures/summary.txt and PRESENTERS when decided, then re-run.
"""
from __future__ import annotations

import re
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.slide import Slide
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "presentation" / "UG_Project_PPT_Template.pptx"
FIGS = ROOT / "figures"
OUT = ROOT / "presentation" / "FYP_Presentation.pptx"

# ----------------------------------------------------------------------------- fill these in
PRESENTERS: dict[str, str] = {}  # slide key -> name, e.g. {"title": "Chanswarang Boro", "intro": ...}; see SLIDE KEYS below

TBD = "TBD"
NUMBERS: dict[str, str] = {
    # single element (CST_GUIDE §1–3) — from figures/summary.txt and the CST parameter list
    "Ws": TBD, "Ls": TBD, "R": TBD, "Lg": TBD, "Wf": "≈ 3", "g": "0.3–0.4", "d": TBD,
    "single_band": TBD,          # e.g. "2.9–11.2 GHz"
    "single_gain": TBD,          # e.g. "2.1–4.8 dBi"
    # AMC (§4–5)
    "p": "10", "b": "9.4", "wr": "0.6", "a": "6",   # starting values; update to the final cell
    "amc_band": TBD,             # ±90° in-phase band, e.g. "5.1–8.3 GHz"
    "h": TBD,                    # chosen air gap in mm
    "amc_gain_delta": TBD,       # e.g. "+2.8" (dB, AMC vs no reflector, band-average or peak — say which in notes)
    "amc_vs_pec": TBD,           # e.g. "+1.5 dB over PEC"
    # MIMO (§6–7)
    "Wb": TBD, "wc": TBD,
    "mimo_band": TBD,            # common −10 dB band of all ports
    "iso": TBD,                  # worst-case isolation, dB
    "ecc": TBD, "ecc_ff": TBD, "dg": TBD, "tarc": TBD, "ccl": TBD, "meg": TBD,
    "mimo_gain_ms": TBD,         # peak realized gain with AMC
    "mesh_delta": TBD,           # max |ΔS11| in band, 15 vs 25 cells/λ (CST_GUIDE §1), e.g. "0.6"
}


def n(key: str) -> str:
    return NUMBERS.get(key, TBD)


# ----------------------------------------------------------------------------- style (template: Times New Roman, white)
FONT = "Times New Roman"
INK = RGBColor(0x00, 0x00, 0x00)
NAVY = RGBColor(0x1F, 0x49, 0x7D)      # template theme dk2
TINT = RGBColor(0xDC, 0xE6, 0xF2)      # light tint of dk2
MUTED = RGBColor(0x59, 0x59, 0x59)
LINE = RGBColor(0xA6, 0xA6, 0xA6)
PEND_FILL = RGBColor(0xF2, 0xF2, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x37, 0x86, 0x3C)

SW, SH = 10.0, 7.5          # 4:3 inches
MX = 0.5                    # side margin
TOP = 1.45                  # content top
BOTTOM = 6.85               # content bottom (footer below)
FOOTER = "Wideband MIMO Antenna with Metasurface  |  Dept. of ECE, NIT Silchar"


def I(x: float) -> Emu:
    return Inches(x)


# ----------------------------------------------------------------------------- text helpers
_TOKEN = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|_\{[^}]*\}|\^\{[^}]*\})")


def add_runs(p, text: str, size: float, color: RGBColor = INK, bold: bool = False, italic: bool = False) -> None:
    """Mini markup: _{sub} ^{sup} **bold** *italic*."""
    for tok in _TOKEN.split(text):
        if not tok:
            continue
        if tok.startswith("**"):
            add_runs(p, tok[2:-2], size, color, True, italic)
            continue
        if tok.startswith("*") and not tok.startswith("*{"):
            add_runs(p, tok[1:-1], size, color, bold, True)
            continue
        r = p.add_run()
        b, it, base = bold, italic, None
        if tok.startswith("_{"):
            tok, base = tok[2:-1], "-25000"
        elif tok.startswith("^{"):
            tok, base = tok[2:-1], "30000"
        r.text = tok
        f = r.font
        f.name, f.size, f.bold, f.italic = FONT, Pt(size), b, it
        f.color.rgb = color
        rPr = r._r.get_or_add_rPr()
        rPr.set("lang", "en-IN")
        if base:
            rPr.set("baseline", base)
        cs = rPr.makeelement(qn("a:cs"), {"typeface": FONT})
        rPr.append(cs)


def set_bullet(p, char: str | None) -> None:
    pPr = p._p.get_or_add_pPr()
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buFont"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    if char is None:
        pPr.set("marL", "0")
        pPr.set("indent", "0")
        pPr.append(pPr.makeelement(qn("a:buNone"), {}))
    else:
        pPr.set("marL", "228600" if p.level == 0 else "457200")
        pPr.set("indent", "-228600")
        pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
        pPr.append(pPr.makeelement(qn("a:buChar"), {"char": char}))


Para = str | tuple  # "text" or (text, {"size":..,"level":..,"bold":..,"color":..,"bullet":..,"align":..,"italic":..})


def fill(tf, paras: list[Para], size: float = 16, bullet: str | None = None, space: float = 4,
         color: RGBColor = INK, align: PP_ALIGN | None = None) -> None:
    tf.clear()
    tf.word_wrap = True
    for k, item in enumerate(paras):
        text, o = (item, {}) if isinstance(item, str) else item
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.level = o.get("level", 0)
        set_bullet(p, o.get("bullet", bullet if p.level == 0 else "–") if (bullet or "bullet" in o) else None)
        p.space_after = Pt(o.get("space", space))
        if o.get("align", align) is not None:
            p.alignment = o.get("align", align)
        add_runs(p, text, o.get("size", size if p.level == 0 else size - 2), o.get("color", color),
                 o.get("bold", False), o.get("italic", False))


def textbox(slide: Slide, x: float, y: float, w: float, h: float, paras: list[Para], size: float = 16,
            bullet: str | None = None, color: RGBColor = INK, align: PP_ALIGN | None = None,
            anchor: MSO_ANCHOR = MSO_ANCHOR.TOP, margin: float = 0.05, name: str | None = None):
    tb = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = I(margin)
    tf.margin_top = tf.margin_bottom = I(0.03)
    tf.vertical_anchor = anchor
    fill(tf, paras, size, bullet, color=color, align=align)
    return tb


def box(slide: Slide, x: float, y: float, w: float, h: float, paras: list[Para] | None = None, size: float = 14,
        fill_rgb: RGBColor = TINT, line: RGBColor | None = None, color: RGBColor = INK,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, align: PP_ALIGN = PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
        bullet: str | None = None, dash: bool = False):
    s = slide.shapes.add_shape(shape, I(x), I(y), I(w), I(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = fill_rgb
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
        if dash:
            s.line.dash_style = 4  # MSO_LINE.DASH
    tf = s.text_frame
    tf.margin_left = tf.margin_right = I(0.08)
    tf.margin_top = tf.margin_bottom = I(0.04)
    tf.vertical_anchor = anchor
    if paras:
        fill(tf, paras, size, bullet, color=color, align=align, space=2)
    return s


def arrow(slide: Slide, x1: float, y1: float, x2: float, y2: float, color: RGBColor = NAVY, width: float = 1.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return c


def table(slide: Slide, x: float, y: float, w: float, rows: list[list[str]], col_w: list[float],
          size: float = 12, row_h: float = 0.3, highlight_last: bool = False):
    shp = slide.shapes.add_table(len(rows), len(rows[0]), I(x), I(y), I(w), I(row_h * len(rows)))
    tbl = shp.table
    tbl.first_row = True
    tbl.horz_banding = False
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = I(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = I(row_h)
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = I(0.05)
            cell.margin_top = cell.margin_bottom = I(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            head = i == 0
            last = highlight_last and i == len(rows) - 1
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if head else (TINT if last else (WHITE if i % 2 else RGBColor(0xF5, 0xF7, 0xFA)))
            fill(cell.text_frame, [val], size, color=WHITE if head else INK,
                 align=PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER, space=0)
            if head or last:
                for r in cell.text_frame.paragraphs[0].runs:
                    r.font.bold = True
    return shp


# ----------------------------------------------------------------------------- figures (real or pending)
PENDING: list[str] = []


def fig_exists(name: str) -> bool:
    return (FIGS / f"{name}.png").exists()


def figure(slide: Slide, name: str, x: float, y: float, w: float, h: float, caption: str | None = None,
           cap_size: float = 12) -> None:
    """Place figures/<name>.png fitted inside (x, y, w, h); otherwise a marked pending box."""
    cap_h = 0.32 if caption else 0.0
    fh = h - cap_h
    path = FIGS / f"{name}.png"
    if path.exists():
        with Image.open(path) as im:
            iw, ih = im.size
        scale = min(w / iw, fh / ih)
        pw, ph = iw * scale, ih * scale
        pic = slide.shapes.add_picture(str(path), I(x + (w - pw) / 2), I(y + (fh - ph) / 2), I(pw), I(ph))
        pic.name = name
    else:
        PENDING.append(name)
        box(slide, x, y, w, fh, [("Result pending — CST simulation", {"bold": True, "size": 13, "color": MUTED}),
                                 (f"figures/{name}.png", {"size": 10, "color": MUTED, "italic": True})],
            fill_rgb=PEND_FILL, line=LINE, dash=True, shape=MSO_SHAPE.RECTANGLE)
    if caption:
        textbox(slide, x, y + fh, w, cap_h, [caption], size=cap_size, color=MUTED, align=PP_ALIGN.CENTER)


# ----------------------------------------------------------------------------- slide frame
def strip_body(slide: Slide) -> None:
    for ph in list(slide.placeholders):
        if ph.placeholder_format.idx != 0:
            ph._element.getparent().remove(ph._element)


def set_title(slide: Slide, text: str, section: str | None, size: float = 26) -> None:
    t = slide.shapes.title
    t.left, t.top, t.width, t.height = I(MX), I(0.42), I(SW - 2 * MX), I(0.95)
    tf = t.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    fill(tf, [text], size, align=PP_ALIGN.CENTER, color=NAVY, space=0)
    for r in tf.paragraphs[0].runs:
        r.font.bold = True
    if section:
        textbox(slide, MX, 0.12, SW - 2 * MX, 0.3, [section.upper()], size=11, color=MUTED,
                align=PP_ALIGN.CENTER, name="Section label")


def footer(slide: Slide, num: int) -> None:
    textbox(slide, MX, 7.05, 7.5, 0.3, [FOOTER], size=10, color=MUTED, name="Footer")
    textbox(slide, SW - MX - 1.0, 7.05, 1.0, 0.3, [str(num)], size=10, color=MUTED, align=PP_ALIGN.RIGHT,
            name="Slide number")


def notes(slide: Slide, key: str, text: str) -> None:
    who = PRESENTERS.get(key, "TBD")
    slide.notes_slide.notes_text_frame.text = f"Presenter: {who}\n\n{text.strip()}"


# ============================================================================= references (IEEE, order of first citation)
REFS: dict[str, str] = {
    "fcc": "Federal Communications Commission, \u201cRevision of Part 15 of the Commission\u2019s rules regarding ultra-wideband transmission systems,\u201d First Report and Order, ET Docket 98-153, FCC 02-48, Apr. 2002 (limits codified in 47 CFR \u00a715.517).",
    "sharawi": "M. S. Sharawi, \u201cPrinted multi-band MIMO antenna systems and their performance metrics,\u201d *IEEE Antennas Propag. Mag.*, vol. 55, no. 5, pp. 218\u2013232, Oct. 2013, doi: 10.1109/MAP.2013.6735522.",
    "wu": "A. Wu, M. Zhao, P. Zhang, and Z. Zhang, \u201cA compact four-port MIMO antenna for UWB applications,\u201d *Sensors*, vol. 22, no. 15, Art. no. 5788, 2022, doi: 10.3390/s22155788.",
    "zhang": "J. Zhang, C. Du, and R. Wang, \u201cDesign of a four-port flexible UWB-MIMO antenna with high isolation for wearable and IoT applications,\u201d *Micromachines*, vol. 13, no. 12, Art. no. 2141, 2022, doi: 10.3390/mi13122141.",
    "yin": "W. Yin, S. Chen, J. Chang, C. Li, and S. K. Khamas, \u201cCPW fed compact UWB 4-element MIMO antenna with high isolation,\u201d *Sensors*, vol. 21, no. 8, Art. no. 2688, 2021, doi: 10.3390/s21082688.",
    "ramanathan": "K. Ramanathan, S. Gopalakrishnan, and T. Chandrakanthan, \u201cMiniaturized dual and quad port MIMO antenna variants featuring elevated diversity performance for UWB and 5G-midband applications,\u201d *Micromachines*, vol. 16, no. 6, Art. no. 716, 2025, doi: 10.3390/mi16060716.",
    "mohanty": "A. Mohanty and S. Sahu, \u201c4-port UWB MIMO antenna with Bluetooth-LTE-WiMax band-rejection and vias-MCP loaded reflector with improved performance,\u201d *AEU \u2013 Int. J. Electron. Commun.*, vol. 144, Art. no. 154065, 2022, doi: 10.1016/j.aeue.2021.154065.",
    "nirmala": "M. Nirmala and N. Deepika Rani, \u201cA spatially selective and electromagnetically tailored low-profile high-gain notched UWB MIMO antenna for 5G and ultra-wideband wireless systems,\u201d *Franklin Open*, vol. 16, Art. no. 100687, 2026, doi: 10.1016/j.fraope.2026.100687.",
    "alekya": "B. Alekya, N. A. Murugan, and B. T. P. Madhav, \u201cArtificial magnetic conductor-integrated high-gain quad-port hexagonal antenna for conformal ultra-wideband applications,\u201d *ETRI J.*, vol. 48, no. 4, pp. 604\u2013619, 2026, doi: 10.4218/etrij.2025-0252.",
    "kumari": "P. Kumari, R. K. Gangwar, and R. K. Chaudhary, \u201cAn aperture-coupled stepped dielectric resonator UWB MIMO antenna with AMC,\u201d *IEEE Antennas Wireless Propag. Lett.*, vol. 21, no. 10, pp. 2040\u20132044, 2022, doi: 10.1109/LAWP.2022.3189694.",
    "douhi": "S. Douhi, Z. Zahriladha, and A. Eddiai, \u201cA high-gain, low-SAR UWB all-textile two-port MIMO antenna based on an AMC structure for wireless body area networks,\u201d *Sci. Rep.*, vol. 16, Art. no. 25405, 2026, doi: 10.1038/s41598-026-45917-z.",
    "azharuddin": "M. Azharuddin and K. Mondal, \u201cFSS-based gain and isolation optimization in a two-element MIMO antenna for ultra-wideband operations,\u201d *S\u0101dhan\u0101*, vol. 51, no. 1, Art. no. 23, 2026, doi: 10.1007/s12046-025-02989-3.",
    "agrawall": "N. P. Agrawall, G. Kumar, and K. P. Ray, \u201cWide-band planar monopole antennas,\u201d *IEEE Trans. Antennas Propag.*, vol. 46, no. 2, pp. 294\u2013295, Feb. 1998, doi: 10.1109/8.660976.",
    "ray": "K. P. Ray, \u201cDesign aspects of printed monopole antennas for ultra-wide band applications,\u201d *Int. J. Antennas Propag.*, vol. 2008, Art. no. 713858, 2008, doi: 10.1155/2008/713858.",
    "sievenpiper": "D. Sievenpiper, L. Zhang, R. F. J. Broas, N. G. Alex\u00f3poulos, and E. Yablonovitch, \u201cHigh-impedance electromagnetic surfaces with a forbidden frequency band,\u201d *IEEE Trans. Microw. Theory Techn.*, vol. 47, no. 11, pp. 2059\u20132074, Nov. 1999, doi: 10.1109/22.798001.",
    "yang": "F. Yang and Y. Rahmat-Samii, \u201cReflection phase characterizations of the EBG ground plane for low profile wire antenna applications,\u201d *IEEE Trans. Antennas Propag.*, vol. 51, no. 10, pp. 2691\u20132703, Oct. 2003, doi: 10.1109/TAP.2003.817559.",
    "blanch": "S. Blanch, J. Romeu, and I. Corbella, \u201cExact representation of antenna system diversity performance from input parameter description,\u201d *Electron. Lett.*, vol. 39, no. 9, pp. 705\u2013707, May 2003, doi: 10.1049/el:20030495.",
    "thaysen": "J. Thaysen and K. B. Jakobsen, \u201cEnvelope correlation in (N,N) MIMO antenna array from scattering parameters,\u201d *Microw. Opt. Technol. Lett.*, vol. 48, no. 5, pp. 832\u2013834, 2006, doi: 10.1002/mop.21490.",
    "manteghi": "M. Manteghi and Y. Rahmat-Samii, \u201cMultiport characteristics of a wideband cavity backed annular patch antenna for multipolarization operations,\u201d *IEEE Trans. Antennas Propag.*, vol. 53, no. 1, pp. 466\u2013474, Jan. 2005, doi: 10.1109/TAP.2004.838794.",
    "chae": "S. H. Chae, S.-K. Oh, and S.-O. Park, \u201cAnalysis of mutual coupling, correlations, and TARC in WiBro MIMO array antenna,\u201d *IEEE Antennas Wireless Propag. Lett.*, vol. 6, pp. 122\u2013125, 2007, doi: 10.1109/LAWP.2007.893109.",
    "taga": "T. Taga, \u201cAnalysis for mean effective gain of mobile antennas in land mobile radio environments,\u201d *IEEE Trans. Veh. Technol.*, vol. 39, no. 2, pp. 117\u2013131, May 1990, doi: 10.1109/25.54228.",
    "simons": "R. N. Simons, *Coplanar Waveguide Circuits, Components, and Systems*. New York, NY, USA: Wiley, 2001, doi: 10.1002/0471224758.",
}


def cite(*keys: str) -> str:
    order = list(REFS)
    nums = sorted(order.index(k) + 1 for k in keys)
    if len(nums) > 2 and nums == list(range(nums[0], nums[-1] + 1)):
        return f"[{nums[0]}]–[{nums[-1]}]"
    return "".join(f"[{x}]" for x in nums)


def lambda_mm(f_ghz: float) -> float:
    return 299.792458 / f_ghz


# ============================================================================= slides
def s_title(s: Slide) -> None:
    t = s.shapes.title
    t.left, t.top, t.width, t.height = I(MX), I(0.35), I(SW - 2 * MX), I(0.6)
    fill(t.text_frame, ["B. Tech. Project Mid-Semester Evaluation"], 22, color=MUTED, align=PP_ALIGN.CENTER)
    textbox(s, MX, 1.0, SW - 2 * MX, 1.2, ["Wideband MIMO Antenna", "with Metasurface"], size=34, color=NAVY,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, name="Project title")
    for para in s.shapes[-1].text_frame.paragraphs:
        para.space_after = Pt(0)
        para.runs[0].font.bold = True
    textbox(s, MX, 2.25, SW - 2 * MX, 0.45,
            ["*A 4-port CPW-fed UWB (3.1–10.6 GHz) MIMO antenna with an AMC reflector*"], size=17,
            align=PP_ALIGN.CENTER, name="Subtitle")
    team = [("Chanswarang Boro", "2314143"), ("Nishit Baishya", "2314088"),
            ("Anushka Dam", "2314115"), ("Sanjana", "2314060")]
    box(s, 0.8, 2.95, 4.0, 2.3,
        [("Presented by", {"bold": True, "color": NAVY, "size": 16, "space": 6})]
        + [(f"{nm}  ({sid})", {"size": 15}) for nm, sid in team],
        fill_rgb=TINT, align=PP_ALIGN.CENTER)
    box(s, 5.2, 2.95, 4.0, 2.3,
        [("Under the guidance of", {"bold": True, "color": NAVY, "size": 16, "space": 6}),
         ("**Dr. Ujjal Chakraborty**", {"size": 15}), ("Associate Professor, ECE", {"size": 14, "space": 10}),
         ("Co-guide", {"bold": True, "color": NAVY, "size": 16, "space": 4}),
         ("**Mr. Sovan Bhattacharya**", {"size": 15}), ("PhD Scholar, ECE", {"size": 14})],
        fill_rgb=TINT, align=PP_ALIGN.CENTER)
    sub = s.placeholders[1]
    sub.left, sub.top, sub.width, sub.height = I(MX), I(5.55), I(SW - 2 * MX), I(1.35)
    fill(sub.text_frame, ["Department of Electronics and Communication Engineering",
                          "**National Institute of Technology Silchar**",
                          "B.Tech 7th Semester, Session 2026–27  |  October 2026"], 16, align=PP_ALIGN.CENTER, space=2)
    notes(s, "title", """
Good morning. We are presenting our B.Tech project, "Wideband MIMO Antenna with Metasurface", under Dr. Ujjal
Chakraborty with Mr. Sovan Bhattacharya as co-guide. In one line: we design a four-port ultra-wideband MIMO
antenna and place a metasurface reflector behind it to raise the gain without losing bandwidth or isolation.
Introduce the four team members. (≈ 15 s)""")


def s_objectives(s: Slide) -> None:
    set_title(s, "Four objectives, each with a measurable target", "Objectives")
    body = s.placeholders[1]
    body.left, body.top, body.width, body.height = I(MX), I(TOP + 0.05), I(4.9), I(5.3)
    fill(body.text_frame, [
        ("**Design** a CPW-fed octagonal monopole that covers the full 3.1–10.6 GHz UWB band", {"bullet": "1"}),
        ("**Optimise** it one parameter at a time (ground length L_{g}, then radius R)", {"bullet": "2"}),
        ("**Add a metasurface**: a dual-resonant AMC reflector behind the antenna, benchmarked against "
         "no reflector and a PEC plate", {"bullet": "3"}),
        ("**Build a 4-port MIMO** (orthogonal elements, common ground) and verify its diversity performance",
         {"bullet": "4"}),
    ], 17, space=12)
    for p in body.text_frame.paragraphs:  # numbered list
        pPr = p._p.get_or_add_pPr()
        for el in pPr.findall(qn("a:buChar")) + pPr.findall(qn("a:buFont")):
            pPr.remove(el)
        pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": FONT}))
        pPr.append(pPr.makeelement(qn("a:buAutoNum"), {"type": "arabicPeriod"}))
    rows = [["Metric", "Target"],
            ["|S_{11}| (all ports)", "≤ −10 dB, 3.1–10.6 GHz"],
            ["Isolation", "> 15 dB (aim 20 dB)"],
            ["Gain increase (AMC)", "+2 to +4 dB vs none & PEC"],
            ["ECC", "< 0.01"],
            ["Diversity gain", "> 9.95 dB"],
            ["TARC", "< −10 dB"],
            ["CCL", "< 0.4 bit/s/Hz"],
            ["MEG", "≈ −3 dB, ports within 3 dB"]]
    textbox(s, 5.6, TOP + 0.05, 3.9, 0.35, ["Design targets"], size=16, color=NAVY)
    s.shapes[-1].text_frame.paragraphs[0].runs[0].font.bold = True
    table(s, 5.6, TOP + 0.45, 3.9, rows, [1.75, 2.15], size=13, row_h=0.42)
    notes(s, "objectives", """
We have four objectives. First, a single CPW-fed octagonal monopole covering the whole UWB band. Second, optimise
it systematically: ground length first, then patch radius. Third, the metasurface: an artificial magnetic
conductor reflector that should add 2 to 4 dB of gain, and we compare it fairly against no reflector and a plain
metal (PEC) plate at the same gap. Fourth, the four-port MIMO with orthogonal elements on a shared ground, judged by
the standard diversity metrics. The targets on the right are typical values from recent four-port UWB papers.""")


def s_methodology(s: Slide) -> None:
    set_title(s, "The design is built and optimised in seven simulation steps", "Proposed Methodology")
    strip_body(s)
    steps = ["1  CPW-fed octagonal monopole", "2  Sweep ground length L_{g}", "3  Sweep patch radius R",
             "4  AMC unit cell (Floquet)", "5  Single + reflector: none / PEC / AMC, gap h",
             "6  4-port orthogonal MIMO", "7  MIMO + AMC", "Diversity metrics: ECC, DG, TARC, CCL, MEG"]
    bw, bh, gap = 2.0, 0.85, (SW - 2 * MX - 4 * 2.0) / 3
    y1, y2 = TOP + 0.1, TOP + 1.45
    xs = [MX + k * (bw + gap) for k in range(4)]
    for k in range(4):
        box(s, xs[k], y1, bw, bh, [steps[k]], size=14)
    for k, xi in enumerate(reversed(xs)):  # row 2 runs right-to-left (snake)
        last = k == 3
        box(s, xi, y2, bw, bh, [steps[4 + k]], size=14 if not last else 13,
            fill_rgb=NAVY if last else TINT, color=WHITE if last else INK)
    for k in range(3):
        arrow(s, xs[k] + bw + 0.03, y1 + bh / 2, xs[k + 1] - 0.03, y1 + bh / 2)
        arrow(s, xs[3 - k] - 0.03, y2 + bh / 2, xs[2 - k] + bw + 0.03, y2 + bh / 2)
    arrow(s, xs[3] + bw / 2, y1 + bh + 0.03, xs[3] + bw / 2, y2 - 0.03)
    cy, cw = TOP + 2.65, (SW - 2 * MX - 0.4) / 3
    cols = [("Tools", ["CST Studio Suite: time-domain solver; Floquet unit-cell solver",
                       "Python + scikit-rf: ECC, DG, TARC, CCL, MEG from exported data"]),
            ("Materials", ["FR-4, ε_{r} = 4.3, tan δ = 0.025, 1.6 mm", "Copper, 35 µm",
                           "Band of interest 2–12 GHz"]),
            ("Method", ["One parameter at a time", "Best = widest share of 3.1–10.6 GHz with |S_{11}| ≤ −10 dB",
                        "Final runs: 20–25 cells/λ mesh"])]
    for k, (head, items) in enumerate(cols):
        x = MX + k * (cw + 0.2)
        box(s, x, cy, cw, BOTTOM - cy, [(head, {"bold": True, "color": NAVY, "size": 16, "bullet": None,
                                                 "align": PP_ALIGN.LEFT, "space": 6})]
            + [(it, {"bullet": "•"}) for it in items],
            size=15, fill_rgb=RGBColor(0xF5, 0xF7, 0xFA), align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
    notes(s, "methodology", """
This is the whole flow. We start from one CPW-fed octagonal monopole and optimise it in two sweeps: ground length,
then radius. "Best" has a precise meaning: the largest fraction of 3.1–10.6 GHz with S11 below −10 dB, with the
worst in-band S11 as the tie-break. In parallel we design the AMC unit cell with a Floquet port. Then we put the
reflector behind the single antenna and compare no reflector, a PEC plate and the AMC at the same gap h. Finally we
build the four-port orthogonal MIMO, with and without the AMC, and evaluate the diversity metrics. All simulations
are in CST. Post-processing is our own Python code, so every figure in this talk comes from the same exported data.""")


def s_geometry(s: Slide) -> None:
    set_title(s, "Single element: a regular octagon fed by a 50 Ω coplanar waveguide", "Proposed Methodology")
    figure(s, "cst_single_geometry", MX, TOP, 4.4, BOTTOM - TOP, "CST front view, parameters labelled")
    rows = [["Parameter", "Meaning", "Value (mm)"],
            ["W_{s} × L_{s}", "Substrate size", f"{n('Ws')} × {n('Ls')}"],
            ["h_{s}", "FR-4 thickness", "1.6"],
            ["R", "Octagon circumradius", n("R")],
            ["L_{g}", "CPW ground length", n("Lg")],
            ["W_{f}", "Feed width", n("Wf")],
            ["g", "CPW slot gap", n("g")],
            ["d", "Ground–patch gap", n("d")]]
    table(s, 5.15, TOP, 4.35, rows, [1.15, 1.95, 1.25], size=13, row_h=0.4)
    textbox(s, 5.15, TOP + 3.45, 4.35, BOTTOM - TOP - 3.45, [
        f"Feed and ground are on one side {cite('simons')}: no via, and the ground can be shared later",
        "Octagon = cylinder with 8 segments: one parameter (R) sets its size",
        "50 Ω checked from the CST port line impedance before any sweep",
    ], size=14, bullet="•")
    notes(s, "geometry", """
The radiator is a regular octagon, drawn as an eight-segment cylinder, so its size is controlled by a single
parameter, the circumradius R. It is fed by a coplanar waveguide. Signal and ground sit on the same face, which
makes the feed easy to fabricate and later lets four elements share one ground. Feed width and slot gap were set so
the port line impedance reads about 50 ohms; we checked that before running any sweep. Substrate: 1.6 mm FR-4.
(Values marked TBD come from the final CST parameter list.)""")


STAGES: list[tuple[str, str | None]] = [
    ("Literature survey and design targets", None),
    ("Python post-processing for all MIMO metrics (tested)", None),
    ("Single element modelled, 50 Ω CPW feed", "fig_single_Lg_sweep"),
    ("L_{g} and R sweeps → optimised element", "fig_single_R_sweep"),
    ("AMC unit cell: reflection phase", "fig_unitcell_reflection"),
    ("Reflector study: none vs PEC vs AMC", "fig_reflector_compare_gain_eff"),
    ("4-port orthogonal MIMO", "fig_mimo_noMS_reflection"),
    ("MIMO + AMC, diversity metrics", "fig_mimo_ecc"),
]


def s_work_done(s: Slide) -> None:
    done = [fig is None or fig_exists(fig) for _, fig in STAGES]
    set_title(s, f"Work done till now: {sum(done)} of {len(STAGES)} project stages complete", "Work Done Till Now")
    strip_body(s)
    textbox(s, MX, TOP, 4.3, 0.35, ["Completed tasks"], size=16, color=NAVY)
    s.shapes[-1].text_frame.paragraphs[0].runs[0].font.bold = True
    for k, ((label, _), ok) in enumerate(zip(STAGES, done)):
        y = TOP + 0.5 + k * 0.6
        c = s.shapes.add_shape(MSO_SHAPE.OVAL, I(MX), I(y + 0.06), I(0.34), I(0.34))
        c.shadow.inherit = False
        c.fill.solid()
        c.fill.fore_color.rgb = GREEN if ok else WHITE
        c.line.color.rgb = GREEN if ok else LINE
        fill(c.text_frame, ["✓" if ok else ""], 12, color=WHITE, align=PP_ALIGN.CENTER)
        c.text_frame.margin_left = c.text_frame.margin_right = 0
        textbox(s, MX + 0.45, y, 3.95, 0.5, [label + ("" if ok else "  *(in progress)*")], size=14,
                color=INK if ok else MUTED, anchor=MSO_ANCHOR.MIDDLE)
    x2 = 5.1
    textbox(s, x2, TOP, SW - MX - x2, 0.35, ["Challenges → how we handle them"], size=16, color=NAVY)
    s.shapes[-1].text_frame.paragraphs[0].runs[0].font.bold = True
    challenges = [
        ("Slow PC: one run takes over 10 min",
         "Field monitors off and coarser mesh during sweeps; 4-fold symmetry → excite port 1 only (≈ 4× faster)"),
        ("Sizing the CPW waveguide port",
         "Port spans feed, both slots and part of the ground; check ≈ 50 Ω line impedance"),
        ("Port extension cuts into the AMC below",
         "Limit the downward extension to h − 0.5 mm, then re-check the impedance"),
        ("No AMC is in-phase over all of UWB",
         "Dual-resonant cell (ring + patch) widens the band; gap h tunes the rest"),
    ]
    y = TOP + 0.5
    for head, fix in challenges:
        box(s, x2, y, SW - MX - x2, 1.12, [(f"**{head}**", {"size": 14, "space": 2}), (fix, {"size": 13})],
            fill_rgb=RGBColor(0xF5, 0xF7, 0xFA), align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
        y += 1.22
    notes(s, "work_done", """
This slide summarises progress. The green ticks update automatically from our result files. The literature survey,
the design targets and the Python post-processing are complete, and the post-processing has unit tests for every
metric formula. The simulation stages are ticked as their CST results come in. The main practical challenge is
simulation time: one run takes more than ten minutes on our PC. So sweeps run with field monitors off and a coarser
mesh, and the MIMO sweeps use the four-fold symmetry: exciting only port 1 gives every S-parameter. Two CST-specific
issues are the CPW port size and the port extension cutting into the AMC; for both we re-check the port
impedance against 50 ohms.""")


def s_sweeps(s: Slide) -> None:
    set_title(s, f"Ground length L_{{g}} sets the low-band match; radius R sets the bandwidth", "Work Done Till Now")
    w = (SW - 2 * MX - 0.3) / 2
    figure(s, "fig_single_Lg_sweep", MX, TOP, w, 3.9, "|S_{11}| vs L_{g} (nominal ± 2 mm, 1 mm steps)")
    figure(s, "fig_single_R_sweep", MX + w + 0.3, TOP, w, 3.9, f"|S_{{11}}| vs R (L_{{g}} fixed at {n('Lg')} mm)")
    box(s, MX, TOP + 4.1, w, BOTTOM - TOP - 4.1,
        [f"**Chosen L_{{g}} = {n('Lg')} mm**", "L_{g} mainly moves the lower band edge and the depth of the first resonance"],
        size=14, fill_rgb=TINT)
    box(s, MX + w + 0.3, TOP + 4.1, w, BOTTOM - TOP - 4.1,
        [f"**Chosen R = {n('R')} mm**", f"Optimised element: −10 dB band {n('single_band')}"],
        size=14, fill_rgb=TINT)
    notes(s, "sweeps", """
Two sweeps, one parameter at a time. Left: ground length. It mostly controls the coupling between the ground edge
and the octagon, so it moves the lower band edge and the impedance match in the low band. We picked the value
that covers the largest share of 3.1–10.6 GHz; our script ranks every curve by that coverage and then by the worst
in-band S11. Right: with Lg fixed, the octagon radius. A larger radius lowers the first resonance, and the right
value keeps all the resonances overlapping across the band. The optimised element's −10 dB band is in the box.
(Gain and efficiency of the optimised element are on backup slide B1.)""")


def s_unit_cell(s: Slide) -> None:
    set_title(s, f"AMC unit cell reflects in phase (±90°) over {n('amc_band')}", "Work Done Till Now")
    figure(s, "cst_unitcell", MX, TOP, 3.3, 2.9, "Unit cell: square ring + patch")
    rows = [["Param", "Meaning", "mm"], ["p", "Period", n("p")], ["b", "Ring outer side", n("b")],
            ["w_{r}", "Ring width", n("wr")], ["a", "Patch side", n("a")]]
    table(s, MX, TOP + 3.05, 3.3, rows, [0.7, 1.8, 0.8], size=13, row_h=0.36)
    figure(s, "fig_unitcell_reflection", MX + 3.6, TOP, SW - 2 * MX - 3.6, 4.2,
           "Reflection phase and magnitude, normal incidence")
    box(s, MX + 3.6, TOP + 4.35, SW - 2 * MX - 3.6, BOTTOM - TOP - 4.35,
        ["Two resonances (ring and patch) merge into one wider in-phase band",
         "Check: with the top metal removed the phase is ≈ ±180°, like a PEC"],
        size=14, fill_rgb=TINT, align=PP_ALIGN.LEFT, bullet="•")
    notes(s, "unit_cell", """
This is the metasurface building block: a square ring with a square patch inside, on grounded 1.6 mm FR-4. We
simulate one cell with periodic (unit-cell) boundaries and a Floquet port de-embedded to the surface. A metal plate
reflects with a 180° phase. The AMC reflects in phase near resonance, and we use the usual ±90° window as its
bandwidth. Ring and patch resonate at two nearby frequencies, which widens that window. As a sanity check, removing
the top metal gives ±180°, exactly like a PEC. Honest limitation: no single passive AMC is in phase over the whole
3.1–10.6 GHz band. The air gap h compensates for part of that, as the next slide shows.""")


def s_reflector(s: Slide) -> None:
    set_title(s, f"The AMC adds {n('amc_gain_delta')} dB of gain and beats a PEC plate at the same gap",
              "Work Done Till Now")
    wl = 4.75
    figure(s, "fig_reflector_compare_gain_eff", MX, TOP, wl, 3.75, "Realized gain & efficiency: none / PEC / AMC")
    wr = SW - 2 * MX - wl - 0.25
    figure(s, "fig_single_gap_sweep", MX + wl + 0.25, TOP, wr, 2.6, "|S_{11}| vs air gap h")
    figure(s, "fig_reflector_compare_s11", MX + wl + 0.25, TOP + 2.75, wr, 2.6, "|S_{11}|: none / PEC / AMC")
    box(s, MX, TOP + 3.9, wl, BOTTOM - TOP - 3.9,
        [f"**Gap h = {n('h')} mm**, same for PEC and AMC",
         f"AMC vs no reflector: {n('amc_gain_delta')} dB",
         f"AMC vs PEC plate: {n('amc_vs_pec')}"],
        size=14, fill_rgb=TINT, align=PP_ALIGN.LEFT, bullet="•")
    notes(s, "reflector", """
This is the key comparison and the answer to "why a metasurface and not a metal plate?". All three cases use
identical mesh and monitors. First the gap sweep at the top right: h trades matching against gain. At the chosen
gap, a PEC plate reflects with 180° phase, so the reflected wave partly cancels the direct one in parts of the band.
The AMC reflects in phase, so the two add up and the back lobe goes forward. We quote the gain increase over the
band where it is at least 2 dB rather than a single best frequency, and we say openly where it drops.""")


def s_mimo(s: Slide) -> None:
    set_title(s, f"4-port orthogonal layout: all ports matched, isolation > {n('iso')} dB", "Work Done Till Now")
    figure(s, "cst_mimo_layout", MX, TOP, 3.9, BOTTOM - TOP - 1.2,
           "Elements rotated 90°, common ground via central plus strip")
    box(s, MX, BOTTOM - 1.05, 3.9, 1.05,
        [f"Board {n('Wb')} × {n('Wb')} mm, strip width w_{{c}} = {n('wc')} mm",
         "Orthogonal elements → polarization diversity"],
        size=13, fill_rgb=TINT, align=PP_ALIGN.LEFT, bullet="•")
    x = MX + 4.15
    w = SW - MX - x
    figure(s, "fig_mimo_noMS_reflection", x, TOP, w, 2.65, "S_{11}–S_{44}")
    figure(s, "fig_mimo_noMS_coupling", x, TOP + 2.7, w, 2.65, "All six couplings S_{ij}")
    notes(s, "mimo", """
The four-port MIMO is made by rotate-copying the optimised element in 90° steps, one element per board edge.
Neighbouring elements are orthogonally polarised, which is what keeps the coupling low on a compact board. All four
CPW grounds are joined through a central plus-shaped strip. A real device has one ground, and keeping the grounds
separate would give unrealistically good isolation. Top right: S11 to S44 against the −10 dB line. Bottom right: all six couplings; the
worst of them is the isolation figure in the title. (Check the title claim against figures/summary.txt.)""")


def s_mimo_ms(s: Slide) -> None:
    set_title(s, f"With the AMC, the MIMO keeps its matching and isolation and gains up to {n('mimo_gain_ms')} dBi",
              "Work Done Till Now")
    wl = 4.6
    figure(s, "fig_mimo_ms_comparison", MX, TOP, wl, 2.65, "S-parameters: without vs with AMC")
    figure(s, "fig_mimo_gain_gain_eff", MX, TOP + 2.7, wl, 2.65, "Realized gain & efficiency, port 1")
    x = MX + wl + 0.25
    w = SW - MX - x
    for k, f in enumerate((4, 7, 10)):
        figure(s, f"fig_pattern_MS_p1_{f}GHz", x, TOP + k * 1.8, w, 1.75, f"Port 1 pattern with AMC, {f} GHz",
               cap_size=11)
    notes(s, "mimo_ms", """
Now the full MIMO with the AMC reflector behind it. Top left: matching and coupling with and without the AMC on the
same axes. The reflector must not cost us bandwidth or isolation, and these curves show whether it does. Bottom
left: realized gain and efficiency for port 1. Right: radiation patterns at 4, 7 and 10 GHz, two principal planes
each. The AMC should lower the back lobe and push the main beam forward. At 10 GHz the pattern is expected to be
less regular because the antenna is electrically larger there. (Describe what the real plots show.) (Port 1 vs port 2 patterns are on backup slide B3.)""")


def s_diversity(s: Slide) -> None:
    set_title(s, f"Diversity metrics: ECC < {n('ecc')}, DG > {n('dg')} dB across the UWB band", "Work Done Till Now")
    w, h, g = (SW - 2 * MX - 0.4) / 3, 2.6, 0.2
    names = [("fig_mimo_ecc", "ECC (from S-parameters)"), ("fig_mimo_dg", "Diversity gain"),
             ("fig_mimo_tarc", "TARC"), ("fig_mimo_ccl", "Channel capacity loss"), ("fig_mimo_meg", "MEG")]
    for k, (nm, cap) in enumerate(names):
        r, c = divmod(k, 3)
        figure(s, nm, MX + c * (w + g), TOP + r * (h + 0.15), w, h, cap, cap_size=11)
    rows = [["Metric", "Target", "Ours"], ["ECC (S)", "< 0.01", n("ecc")], ["ECC (far-field)", "< 0.01", n("ecc_ff")],
            ["DG (dB)", "> 9.95", n("dg")], ["TARC (dB)", "< −10", n("tarc")], ["CCL (b/s/Hz)", "< 0.4", n("ccl")],
            ["MEG (dB)", "≈ −3, Δ < 3", n("meg")]]
    table(s, MX + 2 * (w + g), TOP + h + 0.15, w, rows, [1.2, 0.95, w - 2.15], size=11, row_h=0.355)
    notes(s, "diversity", """
These are the standard MIMO diversity metrics, all computed from the four-port S-parameters with our own script.
Envelope correlation coefficient: below 0.01 means the ports see nearly independent channels. FR-4 is lossy, and the
S-parameter formula assumes a lossless antenna, so we also compute ECC from the far-field patterns; both are in the
table. Diversity gain follows from ECC and should approach 10 dB. TARC accounts for all four ports excited at once
with random phases. Channel capacity loss should stay below 0.4 bit/s/Hz. Mean effective gain should be similar for
all ports, which shows the ports are balanced. Lead with the isolation and gain story: ECC and DG are good in
every paper.""")


def s_future(s: Slide) -> None:
    set_title(s, "Next: fabricate the prototype and validate it by measurement", "Future Work")
    strip_body(s)
    pending = [label for (label, fig) in STAGES if fig and not fig_exists(fig)]
    steps = (["Finish pending CST runs"] if pending else []) + [
        "Fabricate FR-4 antenna and AMC; SMA connectors; spacer for gap h",
        "Measure S-parameters on a VNA (unused ports with 50 Ω loads)",
        "Measure patterns & gain in an anechoic chamber",
        "Compare measured vs simulated; write the paper"]
    bw = (SW - 2 * MX - 0.3 * (len(steps) - 1)) / len(steps)
    for k, st in enumerate(steps):
        x = MX + k * (bw + 0.3)
        c = s.shapes.add_shape(MSO_SHAPE.OVAL, I(x + bw / 2 - 0.3), I(TOP + 0.05), I(0.6), I(0.6))
        c.shadow.inherit = False
        c.fill.solid()
        c.fill.fore_color.rgb = NAVY
        c.line.fill.background()
        fill(c.text_frame, [str(k + 1)], 18, color=WHITE, align=PP_ALIGN.CENTER)
        c.text_frame.paragraphs[0].runs[0].font.bold = True
        box(s, x, TOP + 0.8, bw, 1.75, [st], size=15, fill_rgb=TINT)
        if k:
            arrow(s, x - 0.27, TOP + 0.35, x + bw / 2 - 0.33, TOP + 0.35)
    textbox(s, MX, TOP + 2.85, 9.0, 0.35, ["Anticipated challenges"], size=16, color=NAVY)
    s.shapes[-1].text_frame.paragraphs[0].runs[0].font.bold = True
    textbox(s, MX, TOP + 3.3, 9.0, BOTTOM - TOP - 3.3, [
        "FR-4 permittivity varies between batches (≈ 4.2–4.6), which shifts the resonances. Re-simulate with the measured ε_{r}",
        "The SMA connector and its solder joint are not in the model, so measured S_{11} will differ slightly",
        "Holding the gap h accurately needs non-metallic spacers (foam or nylon)",
        "A 2-port VNA measures a 4-port antenna pair by pair, so the other ports must be terminated",
    ], size=16, bullet="•")
    notes(s, "future", """
Next semester we turn the simulation into hardware. We fabricate the FR-4 antenna and the AMC board, add SMA
connectors and hold the gap with non-metallic spacers. Then we measure S-parameters on a VNA, pair by pair with the
unused ports terminated, and measure patterns and gain in the anechoic chamber. Last, we compare measurement with
simulation and write it up. The main risks are on the slide. FR-4 permittivity varies, so we will re-simulate with
the measured value, and the connector is not in our model, so small differences are expected.""")


def s_conclusion(s: Slide) -> None:
    set_title(s, "Conclusion: a compact 4-port UWB MIMO whose AMC adds gain without losing diversity", "Conclusion")
    strip_body(s)
    stats = [(n("mimo_band"), "−10 dB band, all ports"), (f"> {n('iso')} dB", "port isolation"),
             (f"< {n('ecc')}", "ECC"), (f"{n('amc_gain_delta')} dB", "gain increase from AMC")]
    w = (SW - 2 * MX - 0.3 * 3) / 4
    for k, (val, lab) in enumerate(stats):
        box(s, MX + k * (w + 0.3), TOP + 0.05, w, 1.45,
            [(val, {"size": 22, "bold": True, "color": NAVY, "space": 4}), (lab, {"size": 13})], fill_rgb=TINT)
    textbox(s, MX, TOP + 1.75, 9.0, 0.35, ["Summary of progress"], size=16, color=NAVY)
    s.shapes[-1].text_frame.paragraphs[0].runs[0].font.bold = True
    textbox(s, MX, TOP + 2.15, 9.0, 2.0, [
        "A CPW-fed octagonal monopole was optimised (L_{g}, then R) to cover the UWB band",
        "A dual-resonant AMC reflector raises gain and outperforms a PEC plate at the same gap",
        "The 4-port orthogonal MIMO with a common ground meets the isolation and diversity targets",
    ], size=15, bullet="•")
    textbox(s, MX, TOP + 3.65, 9.0, 0.35, ["Next steps"], size=16, color=NAVY)
    s.shapes[-1].text_frame.paragraphs[0].runs[0].font.bold = True
    textbox(s, MX, TOP + 4.05, 9.0, BOTTOM - TOP - 4.05, [
        "Fabrication → VNA and anechoic-chamber measurement → comparison with simulation → paper"], size=15, bullet="•")
    notes(s, "conclusion", """
To conclude: the four numbers at the top are our headline results, covering bandwidth, isolation, correlation and
the gain added by the metasurface. We optimised a CPW-fed octagonal monopole, designed a dual-resonant AMC that
beats a metal plate at the same gap, and built a four-port orthogonal MIMO on a shared ground that meets the
diversity targets. Next comes fabrication and measurement. Thank you, we welcome your questions.
(Before presenting: confirm every number and claim on this slide against figures/summary.txt.)""")


def s_backup_single(s: Slide) -> None:
    set_title(s, f"B1 · Optimised single element: −10 dB band {n('single_band')}", "Backup")
    w = (SW - 2 * MX - 0.3) / 2
    figure(s, "fig_single_final_s11", MX, TOP, w, 4.5, "|S_{11}|, optimised element (fine mesh)")
    figure(s, "fig_single_final_gain_eff", MX + w + 0.3, TOP, w, 4.5, "Realized gain & total efficiency")
    textbox(s, MX, TOP + 4.7, 9.0, BOTTOM - TOP - 4.7,
            [f"Realized gain {n('single_gain')} over 3.1–10.6 GHz. Mesh check: |S_{{11}}| at 15 vs 25 cells/λ differs by "
             f"≤ {n('mesh_delta')} dB in band"],
            size=14, align=PP_ALIGN.CENTER)
    notes(s, "backup", "Backup for questions about the single element and mesh convergence.")


def s_backup_current(s: Slide) -> None:
    set_title(s, "B2 · Port 1 excited: current stays on element 1", "Backup")
    w = (SW - 2 * MX - 0.4) / 3
    for k, f in enumerate((4, 7, 10)):
        figure(s, f"cst_current_{f}GHz", MX + k * (w + 0.2), TOP, w, 4.6, f"Surface current, {f} GHz")
    textbox(s, MX, TOP + 4.8, 9.0, BOTTOM - TOP - 4.8,
            ["Same fixed colour scale at all three frequencies; ports 2–4 matched (50 Ω)"], size=14,
            align=PP_ALIGN.CENTER)
    notes(s, "backup", """
Backup for coupling questions. With only port 1 excited, the current concentrates on element 1 and its feed. Weak
current on the other elements is direct evidence of low coupling. The central ground strip carries part of the
current, which is how it helps the isolation.""")


def s_backup_patterns(s: Slide) -> None:
    set_title(s, "B3 · Ports 1 and 2 radiate orthogonally (with AMC)", "Backup")
    w = (SW - 2 * MX - 0.3) / 2
    for r, f in enumerate((4, 7, 10)):
        for c, p in enumerate((1, 2)):
            figure(s, f"fig_pattern_MS_p{p}_{f}GHz", MX + c * (w + 0.3), TOP + r * 1.8, w, 1.75,
                   f"Port {p}, {f} GHz", cap_size=11)
    notes(s, "backup", "Backup for pattern-diversity questions: port 2 is the 90°-rotated copy of port 1.")


# ============================================================================= intro & theory
FR4 = RGBColor(0x6B, 0x8E, 0x23)
COPPER = RGBColor(0xC0, 0x7A, 0x3A)


def _board(s: Slide, x: float, y: float, w: float) -> None:
    b = box(s, x, y, w, 0.12, fill_rgb=FR4, shape=MSO_SHAPE.RECTANGLE)
    b.name = "Antenna board"
    box(s, x + w / 2 - 0.35, y - 0.05, 0.7, 0.05, fill_rgb=COPPER, shape=MSO_SHAPE.RECTANGLE)


def _amc(s: Slide, x: float, y: float, w: float) -> None:
    box(s, x, y, w, 0.1, fill_rgb=FR4, shape=MSO_SHAPE.RECTANGLE).name = "AMC substrate"
    box(s, x, y + 0.1, w, 0.04, fill_rgb=COPPER, shape=MSO_SHAPE.RECTANGLE)
    k, cw = 9, w / 9
    for i in range(k):
        box(s, x + i * cw + cw * 0.15, y - 0.04, cw * 0.7, 0.04, fill_rgb=COPPER, shape=MSO_SHAPE.RECTANGLE)


def s_intro(s: Slide) -> None:
    set_title(s, "Compact UWB MIMO boards face two problems: port coupling and low monopole gain",
              "Introduction & Problem Statement")
    body = s.placeholders[1]
    body.left, body.top, body.width, body.height = I(MX), I(TOP), I(5.2), I(BOTTOM - TOP)
    fill(body.text_frame, [
        ("Background", {"bold": True, "color": NAVY, "bullet": None, "space": 2}),
        (f"UWB uses 3.1–10.6 GHz at very low power (−41.3 dBm/MHz) {cite('fcc')}: high data rate over short range",
         {"bullet": "•"}),
        (f"MIMO: several antennas on one device raise capacity without extra bandwidth or power {cite('sharawi')}",
         {"bullet": "•", "space": 8}),
        ("Problem", {"bold": True, "color": NAVY, "bullet": None, "space": 2}),
        ("Four antennas on a small board couple to each other, which correlates their signals", {"bullet": "•"}),
        (f"A printed monopole radiates both ways: reported 4-port UWB designs reach only ≈ 1–7 dBi {cite('yin', 'wu', 'zhang', 'ramanathan')}",
         {"bullet": "•", "space": 8}),
        ("Problem definition", {"bold": True, "color": NAVY, "bullet": None, "space": 2}),
        ("Design a compact 4-port UWB MIMO antenna with high isolation, and use a metasurface reflector to raise its "
         "gain without losing bandwidth or isolation", {"bullet": "•"}),
    ], 15, space=3)
    x, w = 6.05, 3.3
    textbox(s, x - 0.2, TOP, w + 0.4, 0.3, ["Printed monopole alone"], size=13, color=MUTED, align=PP_ALIGN.CENTER)
    _board(s, x, TOP + 1.05, w)
    arrow(s, x + w / 2, TOP + 0.95, x + w / 2, TOP + 0.4, width=2.5)
    arrow(s, x + w / 2, TOP + 1.25, x + w / 2, TOP + 1.8, color=RGBColor(0xC0, 0x50, 0x4D), width=2.5)
    textbox(s, x - 0.2, TOP + 1.85, w + 0.4, 0.3, ["Half the power goes backwards"], size=12, color=MUTED,
            align=PP_ALIGN.CENTER)
    y0 = TOP + 2.55
    textbox(s, x - 0.2, y0, w + 0.4, 0.3, ["With AMC reflector at gap h"], size=13, color=MUTED, align=PP_ALIGN.CENTER)
    _board(s, x, y0 + 1.05, w)
    _amc(s, x, y0 + 2.05, w)
    arrow(s, x + w / 2 - 0.35, y0 + 0.95, x + w / 2 - 0.35, y0 + 0.4, width=2.5)
    arrow(s, x + w / 2 + 0.35, y0 + 1.25, x + w / 2 + 0.35, y0 + 1.95, color=RGBColor(0xC0, 0x50, 0x4D), width=2.5)
    arrow(s, x + w / 2 + 0.6, y0 + 1.95, x + w / 2 + 0.6, y0 + 0.4, width=2.5)
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x + 0.2), I(y0 + 1.2), I(x + 0.2), I(y0 + 2.0))
    c.line.color.rgb = MUTED
    textbox(s, x + 0.25, y0 + 1.4, 0.4, 0.3, ["*h*"], size=14, color=MUTED)
    textbox(s, x - 0.2, y0 + 2.3, w + 0.4, 0.5, ["In-phase reflection adds to the forward beam"], size=12,
            color=MUTED, align=PP_ALIGN.CENTER)
    notes(s, "intro", f"""
Ultra-wideband is the 3.1 to 10.6 GHz band the FCC opened in 2002 for very low-power, high-data-rate links
{cite('fcc')}. MIMO puts several antennas on one device, so capacity grows without more spectrum or power
{cite('sharawi')}. Putting four wideband antennas on a small board causes two problems. First, they couple to each
other, which correlates their signals and cancels the MIMO benefit. Second, a printed monopole radiates equally
forwards and backwards (top sketch), so its gain is low: roughly 1 to 7 dBi, about 4 dBi on average, in recent
four-port designs {cite('yin', 'wu', 'zhang', 'ramanathan')}. The bottom sketch is our idea: a metasurface behind the antenna reflects the backward wave
in phase, so it adds to the forward beam. Our problem is to do that across the whole UWB band without losing
isolation.""")


def s_theory(s: Slide) -> None:
    set_title(s, "Three textbook results set the antenna, the reflector and the metrics", "Proposed Methodology")
    rows = [
        ("Printed UWB monopole", [
            "f_{L} ≈ 7.2 / (L + r + p)  GHz   (lengths in cm)",
            ("Equivalent cylinder: 2πrL = patch area; p = feed gap", {"size": 13, "color": MUTED})],
         [f"A planar monopole behaves like a thick cylinder, so it is wideband {cite('agrawall', 'ray')}",
          "Our octagon estimate: L + r ≈ 2.1R, so f_{L} = 3.1 GHz needs R ≈ 11 mm before substrate loading"]),
        ("AMC reflector", [
            "Z_{s} = jωL / (1 − ω^{2}LC),   f_{0} = 1 / (2π√(LC))",
            ("In-phase band: phase within ±90°; BW ≈ √(L/C)/η", {"size": 13, "color": MUTED})],
         [f"A PEC reflects at 180° and needs λ/4 spacing (24 mm at 3.1 GHz) {cite('sievenpiper')}",
          f"An AMC reflects at 0° near f_{{0}}, so it can sit much closer {cite('sievenpiper', 'yang')}"]),
        ("MIMO diversity", [
            "ECC_{ij} = |Σ_{n} S_{ni}^{∗}S_{nj}|^{2} / [(1 − Σ_{n}|S_{ni}|^{2})(1 − Σ_{n}|S_{nj}|^{2})]",
            ("DG = 10√(1 − ECC^{2});  TARC = √(Σ|b_{i}|^{2}) / √(Σ|a_{i}|^{2})", {"size": 13}),
            ("CCL = −log_{2} det(ψ^{R});  MEG_{i} = 0.5(1 − Σ_{j}|S_{ij}|^{2})", {"size": 13})],
         [f"S-parameter ECC assumes a lossless antenna {cite('blanch', 'thaysen')}, so we also compute it from "
          f"far-field patterns {cite('sharawi')}",
          f"TARC {cite('manteghi')}, CCL {cite('chae')}, MEG {cite('taga')}: lossless port → MEG = −3 dB"]),
    ]
    rh, gap = 1.68, 0.15
    for k, (head, eqs, facts) in enumerate(rows):
        y = TOP + k * (rh + gap)
        box(s, MX, y, 1.55, rh, [head], size=15, fill_rgb=NAVY, color=WHITE)
        s.shapes[-1].text_frame.paragraphs[0].runs[0].font.bold = True
        box(s, MX + 1.65, y, 4.15, rh, [(e, {"space": 3}) if isinstance(e, str) else e for e in eqs], size=14,
            fill_rgb=TINT, align=PP_ALIGN.CENTER)
        textbox(s, MX + 5.95, y, SW - MX - (MX + 5.95), rh, facts, size=13, bullet="•", anchor=MSO_ANCHOR.MIDDLE)
    notes(s, "theory", f"""
Three standard results drive the design. Top: a planar monopole behaves like a thick cylindrical monopole, which is
why it covers a wide band. Its lower band edge is roughly 7.2 over (L + r + p) GHz, with lengths in centimetres
{cite('agrawall', 'ray')}. Applied to our octagon, that gives a starting radius of about 11 mm. The CST sweeps then
refine it, because the FR-4 substrate lowers the frequency further. Middle: an AMC behaves like a parallel LC
resonator {cite('sievenpiper')}. At resonance its surface impedance is very high and it reflects in phase, while a
metal plate reflects at 180° and would need a quarter-wavelength gap, 24 mm at 3.1 GHz. We define the AMC band
where the phase stays within ±90°. Bottom: the MIMO metrics we compute. ECC from S-parameters assumes a lossless
antenna, and FR-4 is lossy, so we also compute it from the far-field patterns {cite('sharawi')}.""")


# ============================================================================= literature (verified: LITERATURE.md)
def s_literature(s: Slide) -> None:
    set_title(s, "Orthogonal layouts already give good isolation; reflectors add gain but lack a fair baseline",
              "Literature Review")
    strip_body(s)
    cards = [
        (f"4-port + reflector {cite('mohanty', 'nirmala', 'alekya')}", [
            f"Metal reflector with via probes: isolation > 15 → > 20 dB, gain 7–9.57 dBi at h ≈ 10 mm {cite('mohanty')}",
            f"Ring FSS: gain 4.7 → 7 dBi {cite('nirmala')}; split-ring AMC: up to 9.6 dBi {cite('alekya')}"]),
        (f"4-port, no reflector {cite('yin', 'wu', 'zhang', 'ramanathan')}", [
            "Orthogonal elements alone give isolation > 15–22 dB and ECC < 0.01–0.08",
            "Peak gain only ≈ 1–7 dBi (≈ 4 dBi average): a monopole radiates on both sides"]),
        (f"2-port + AMC / FSS {cite('kumari', 'douhi', 'azharuddin')}", [
            f"AMC or FSS backing raises gain to ≈ 10–14 dBi {cite('douhi', 'azharuddin')}",
            f"Dual-AMC reflector keeps gain within 4.2–6.5 dBi across the band {cite('kumari')}"]),
    ]
    w = (SW - 2 * MX - 0.4) / 3
    for k, (head, items) in enumerate(cards):
        box(s, MX + k * (w + 0.2), TOP + 0.05, w, 3.0,
            [(head, {"bold": True, "color": NAVY, "size": 16, "space": 8})]
            + [(it, {"bullet": "•", "space": 8}) for it in items],
            size=15, fill_rgb=RGBColor(0xF5, 0xF7, 0xFA), align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
    box(s, MX, TOP + 3.3, SW - 2 * MX, BOTTOM - TOP - 3.3, [
        ("**Research gap**", {"size": 16, "space": 4}),
        (f"In the accessible text of the 4-port reflector papers, only {cite('mohanty')} states the gap h, and none "
         "compares against a PEC plate. None of the ten papers uses CPW-fed octagonal monopoles with a "
         "dual-resonant AMC. **We report gain with vs without the AMC across the band, a PEC baseline, h in mm "
         "and λ, and both S-parameter and far-field ECC.**", {"size": 15}),
    ], fill_rgb=NAVY, color=WHITE, align=PP_ALIGN.LEFT)
    notes(s, "literature", f"""
We reviewed ten recent papers in three groups. The four-port UWB designs without a reflector
{cite('yin', 'wu', 'zhang', 'ramanathan')} show that orthogonal placement alone gives 17 to 22 dB isolation and very
low ECC, so isolation is a solved problem. Their gain is only about 4 dBi on average, because a printed monopole
radiates equally forwards and backwards. Reflectors, FSS and AMC layers fix that {cite('mohanty', 'nirmala', 'alekya',
'douhi', 'azharuddin')}, reaching 7 to 14 dBi. However, the gap to the reflector is often not stated, and no
four-port paper we could read compares the AMC with a plain metal plate. That is our gap: CPW-fed octagons with a
dual-resonant AMC, reported fairly against no reflector and a PEC plate. (Several of these papers were readable only
as abstracts; LITERATURE.md marks which.)""")


def s_outcomes(s: Slide) -> None:
    set_title(s, "Expected outcome: diversity on par with the best 4-port designs, plus AMC gain", "Expected Outcomes")
    strip_body(s)
    deliver = ["Optimised CST models: single element, AMC cell, 4-port with and without AMC",
               "Full metric set: S-parameters, gain, efficiency, patterns, ECC, DG, TARC, CCL, MEG",
               "Fabricated prototype, measurements and a paper draft"]
    w = (SW - 2 * MX - 0.4) / 3
    for k, d in enumerate(deliver):
        box(s, MX + k * (w + 0.2), TOP, w, 0.95, [d], size=13, fill_rgb=TINT)
    rows = [["Ref.", "Ports", "Size (mm)", "Band (GHz)", "Iso. (dB)", "ECC", "Gain (dBi)", "Reflector"],
            [cite("mohanty"), "4", "≈ 0.46λ₀ sq.", "2.08–10.4", "> 20", "—", "7–9.57", "Metal + via probes, h ≈ 10 mm"],
            [cite("nirmala"), "4", "—", "3.1–10.6", "—", "—", "4.7 → 7", "Circular-ring FSS"],
            [cite("alekya"), "4", "60×60×0.1", "2–12", "> 25", "—", "9.6", "Split-ring AMC"],
            [cite("douhi"), "2", "120×80×1", "3.7–14", "> 20", "n/r", "4.33 → 14.26", "AMC, 5 mm foam"],
            [cite("azharuddin"), "2", "20×28×1.6", "3.6–10.6", "> 20", "< 0.005", "9.85", "FSS"],
            [cite("yin"), "4", "38×38×1.6", "3–20", "> 17", "< 0.08", "1.3–6.2", "None"],
            [cite("wu"), "4", "45×45×1.6", "3.1–13.1", "> 17", "< 0.02", "1.6–7.3", "None"],
            [cite("zhang"), "4", "65×65×0.1", "2.9–10.86", "> 22", "< 0.01", "0.95–6.49", "None"],
            [cite("ramanathan"), "4", "60×41×1.6", "2.6–10.8", "> 15", "0.014 avg", "2.3–6.12", "None"],
            ["This work", "4", f"{n('Wb')}×{n('Wb')}×1.6", n("mimo_band"), f"> {n('iso')}", n("ecc"),
             n("mimo_gain_ms"), f"Dual-resonant AMC, h = {n('h')} mm"]]
    rows[1:-1] = sorted(rows[1:-1], key=lambda r: int(r[0].strip("[]")))
    table(s, MX, TOP + 1.15, SW - 2 * MX, rows, [0.82, 0.55, 1.2, 1.05, 0.75, 0.9, 1.18, 2.55], size=11, row_h=0.33,
          highlight_last=True)
    textbox(s, MX, TOP + 1.15 + 0.33 * len(rows) + 0.08, SW - 2 * MX, 0.5,
            ["— = not in the accessible text (abstract only); n/r = not reported. Gain: without → with reflector. "
             "Details and sources in LITERATURE.md"], size=10, color=MUTED)
    notes(s, "outcomes", f"""
The deliverables are on top: verified CST models, the full metric set, and next semester a fabricated, measured
prototype. The table puts our target row next to the literature. Isolation and ECC should match the orthogonal
designs {cite('yin', 'wu', 'zhang', 'ramanathan')}. The gain should be clearly above their ≈ 4 dBi average because of
the AMC, and we report the gap and the PEC baseline that most reflector papers leave out. Dashes mean we could only
read the abstract, so we did not fill numbers we could not verify.""")


REF_PAGES = 2


def s_references(s: Slide, page: int) -> None:
    keys = list(REFS)
    per = -(-len(keys) // REF_PAGES)
    chunk = keys[page * per:(page + 1) * per]
    set_title(s, f"References ({page + 1}/{REF_PAGES})", "Conclusion & References")
    textbox(s, MX, TOP - 0.1, SW - 2 * MX, BOTTOM - TOP + 0.1,
            [(f"[{keys.index(k) + 1}]  {REFS[k]}", {"space": 4}) for k in chunk], size=11)
    notes(s, "references", "Reference list in IEEE style, numbered in order of first citation. Not presented aloud.")


# ============================================================================= assembly
TITLE_ONLY, TITLE_CONTENT = 5, 1
# (key, template slide number or None for a new slide, builder). Template sections keep their order.
SLIDES = [
    ("title", 1, s_title),
    ("intro", 2, s_intro),
    ("objectives", 3, s_objectives),
    ("literature", 4, s_literature),
    ("methodology", 5, s_methodology),
    ("theory", None, s_theory),
    ("geometry", None, s_geometry),
    ("work_done", 6, s_work_done),
    ("sweeps", None, s_sweeps),
    ("unit_cell", None, s_unit_cell),
    ("reflector", None, s_reflector),
    ("mimo", None, s_mimo),
    ("mimo_ms", None, s_mimo_ms),
    ("diversity", None, s_diversity),
    ("future", 7, s_future),
    ("outcomes", 8, s_outcomes),
    ("conclusion", 9, s_conclusion),
    ("references", None, lambda s: s_references(s, 0)),
    ("references_2", None, lambda s: s_references(s, 1)),
    ("backup_single", None, s_backup_single),
    ("backup_current", None, s_backup_current),
    ("backup_patterns", None, s_backup_patterns),
]


def build() -> Path:
    prs = Presentation(str(TEMPLATE))
    template_slides = list(prs.slides)
    ordered: list[tuple[str, Slide, object]] = []
    for key, tnum, fn in SLIDES:
        s = template_slides[tnum - 1] if tnum else prs.slides.add_slide(prs.slide_layouts[TITLE_ONLY])
        ordered.append((key, s, fn))
    id_list = prs.slides._sldIdLst
    by_id = {int(el.get("id")): el for el in id_list}
    wanted = [by_id[s.slide_id] for _, s, _ in ordered]
    for el in list(id_list):
        id_list.remove(el)
    for el in wanted:
        id_list.append(el)
    for num, (key, s, fn) in enumerate(ordered, 1):
        fn(s)
        if num > 1:
            footer(s, num)
    prs.core_properties.title = "Wideband MIMO Antenna with Metasurface"
    prs.core_properties.author = "C. Boro, N. Baishya, A. Dam, Sanjana"
    prs.save(str(OUT))
    return OUT


if __name__ == "__main__":
    out = build()
    print(f"Wrote {out.relative_to(ROOT)}")
    print(f"Pending figures ({len(set(PENDING))}):", ", ".join(dict.fromkeys(PENDING)) or "none")
    print("TBD numbers:", ", ".join(k for k, v in NUMBERS.items() if v == TBD) or "none")
